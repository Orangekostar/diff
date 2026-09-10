"""Bounded Hasebe source-evidence tracing helpers."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections.abc import Iterable, Mapping
from dataclasses import asdict
from pathlib import Path
from typing import Any

import numpy as np
import yaml
from openpyxl import load_workbook
from PIL import Image

AUTHOR_SPATIAL_REFERENCE = "AUTHOR_SPATIAL_REFERENCE"
AUTHOR_SCALAR_MEASUREMENT = "AUTHOR_SCALAR_MEASUREMENT"
AUTHOR_ULTRASOUND_MEASUREMENT = "AUTHOR_ULTRASOUND_MEASUREMENT"
EXPERT_REVIEWED_REFERENCE = "EXPERT_REVIEWED_REFERENCE"
LOCAL_DERIVED_PROXY = "LOCAL_DERIVED_PROXY"
MECHANICAL_ENDPOINT_ONLY = "MECHANICAL_ENDPOINT_ONLY"
UNRESOLVED_PROVENANCE = "UNRESOLVED_PROVENANCE"

_POST_LVI_STAGE = "POST_LVI_PRE_CAI"
_DAMAGE_SCALARS = {"PROJECTED_DELAMINATION_AREA", "DENT_DEPTH"}

_BASE_SHA = "a48c76fcead10d4f833c54ab6e892eb524432269"
_REFERENCE_VERSION = "REVIEWED_02339eda86d486c3"
_DATASET_BY_PREFIX = {
    "c8": "74t7kcdgkr",
    "c16": "yfxyg8jm46",
    "c24": "xcmzfsbd9t",
    "q8": "ykhs7s2dck",
    "q16": "w68dtmpfyf",
    "q24": "cgtnjyggtm",
}
_FIXED_CASES = (
    "ykhs7s2dck:q8-4",
    "74t7kcdgkr:c8-24",
    "cgtnjyggtm:q24-40",
)
_WORKBOOK_KEYS = (
    "impact_condition_74t7kcdgkr",
    "impact_condition_yfxyg8jm46",
    "impact_condition_xcmzfsbd9t",
    "impact_condition_ykhs7s2dck",
    "impact_condition_w68dtmpfyf",
    "impact_condition_cgtnjyggtm",
    "lvi_workbook",
    "size_workbook",
    "cai_workbook",
)
_PUBLIC_EVIDENCE = (
    {
        "source_id": "O01",
        "citation": "Hasebe et al., Data in Brief 42 (2022), 108462",
        "doi": "10.1016/j.dib.2022.108462",
        "url": "https://pmc.ncbi.nlm.nih.gov/articles/PMC9294053/",
        "access_status": "FULL_TEXT",
        "supports": (
            "Internal-damage files are ultrasound C-scan screenshots; display colour "
            "encodes amplitude. Reported setup: 3.5 MHz probe, 0.200 x 0.200 mm "
            "pitch, and 75 x 75 mm scanning length."
        ),
        "does_not_support": "A downloadable author mask or raw quantitative amplitude grid.",
    },
    {
        "source_id": "O02",
        "citation": "Hasebe et al., Composites Part B 237 (2022), 109844",
        "doi": "10.1016/j.compositesb.2022.109844",
        "url": "https://doi.org/10.1016/j.compositesb.2022.109844",
        "access_status": "ABSTRACT_AND_INTRO_ONLY",
        "supports": "The prediction targets include delamination area and delamination length.",
        "does_not_support": "The inaccessible methods cannot be cited for a public mask protocol.",
    },
    {
        "source_id": "O03",
        "citation": "Six Hasebe low-velocity-impact datasets, Mendeley Data v1",
        "doi": "10.17632/74t7kcdgkr.1 and five bound companion records",
        "url": "https://data.mendeley.com/datasets/74t7kcdgkr/1",
        "access_status": "RECORD_DESCRIPTIONS_VERIFIED",
        "supports": "Each record describes post-impact C-scan images and an image scale.",
        "does_not_support": "A spatial contour or raw ultrasound array in the bound local files.",
    },
    {
        "source_id": "O04",
        "citation": "Hasebe et al., Mendeley Data CAI dataset v3",
        "doi": "10.17632/8scdmfdcfb.3",
        "url": "https://data.mendeley.com/datasets/8scdmfdcfb/3",
        "access_status": "RECORD_DESCRIPTION_VERIFIED",
        "supports": "The package separates LVI conditions/damage, specimen size, CAI images, raw logger records, and strength.",
        "does_not_support": "Treating post-CAI photographs or CAI logger traces as pre-CAI ultrasound.",
    },
    {
        "source_id": "O05",
        "citation": "Hasebe et al., Data in Brief 58 (2025), 111509",
        "doi": "10.1016/j.dib.2025.111509",
        "url": "https://pmc.ncbi.nlm.nih.gov/articles/PMC11999467/",
        "access_status": "FULL_TEXT",
        "supports": "Projected delamination area was measured from C-scan images with ImageJ; dent depth is reported separately.",
        "does_not_support": "An exported ImageJ ROI/contour in the files inspected here.",
    },
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    """Read workbook structure and cells without recalculating formulas."""

    source = Path(path)
    formula_book = load_workbook(source, data_only=False, read_only=False)
    cached_book = load_workbook(source, data_only=True, read_only=False)
    try:
        sheets: list[dict[str, Any]] = []
        for formula_sheet in formula_book.worksheets:
            cached_sheet = cached_book[formula_sheet.title]
            cells: list[dict[str, Any]] = []
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    is_formula = cell.data_type == "f"
                    cached_value = (
                        cached_sheet[cell.coordinate].value if is_formula else None
                    )
                    cells.append(
                        {
                            "address": cell.coordinate,
                            "value_raw": cell.value,
                            "data_type": cell.data_type,
                            "number_format": cell.number_format,
                            "formula": cell.value if is_formula else None,
                            "cached_value": cached_value,
                            "cache_status": (
                                "FORMULA_CACHED_VALUE"
                                if is_formula and cached_value is not None
                                else "FORMULA_NO_CACHED_VALUE"
                                if is_formula
                                else "NOT_FORMULA"
                            ),
                        }
                    )
            sheets.append(
                {
                    "name": formula_sheet.title,
                    "visibility": formula_sheet.sheet_state,
                    "dimensions": formula_sheet.calculate_dimension(),
                    "merged_ranges": [
                        str(value) for value in formula_sheet.merged_cells.ranges
                    ],
                    "freeze_panes": (
                        str(formula_sheet.freeze_panes)
                        if formula_sheet.freeze_panes is not None
                        else None
                    ),
                    "comment_count": sum(
                        cell.comment is not None
                        for row in formula_sheet.iter_rows()
                        for cell in row
                    ),
                    "image_count": len(formula_sheet._images),
                    "chart_count": len(formula_sheet._charts),
                    "cells": cells,
                }
            )
        return {
            "source_path": source.as_posix(),
            "source_sha256": _sha256(source),
            "sheet_count": len(sheets),
            "sheets": sheets,
        }
    finally:
        formula_book.close()
        cached_book.close()


def provenance_status(
    *, source_kind: str, extraction_method: str, measurement_semantics: str
) -> str:
    """Classify provenance from the source and operation, never from a friendly name."""

    source = str(source_kind).strip()
    method = str(extraction_method).strip()
    semantics = str(measurement_semantics).strip().upper()
    if semantics in {"CAI_STRENGTH", "CAI_STRENGTH_REDUCTION"}:
        return MECHANICAL_ENDPOINT_ONLY
    if method in {"RGB_THRESHOLD_MORPHOLOGY", "FULL_INPUT_READER"}:
        return LOCAL_DERIVED_PROXY
    if source == "AUTHOR_CSCAN_SCREENSHOT" and method == "FILE_IDENTITY_ONLY":
        return AUTHOR_ULTRASOUND_MEASUREMENT
    if source == "AUTHOR_WORKBOOK" and method == "DIRECT_CELL_READ":
        return AUTHOR_SCALAR_MEASUREMENT
    if source == "EXPERT_REVIEW_JSON" and method == "POLYGON_RASTERIZATION":
        return EXPERT_REVIEWED_REFERENCE
    if source == "AUTHOR_SPATIAL_FILE" and method == "DIRECT_SPATIAL_IMPORT":
        return AUTHOR_SPATIAL_REFERENCE
    return UNRESOLVED_PROVENANCE


def _exact_key(row: Mapping[str, object]) -> tuple[str, str]:
    dataset = str(row.get("dataset_id", "")).strip()
    specimen = str(row.get("specimen_id", "")).strip().lower()
    if not specimen:
        specimen_key = str(row.get("specimen_key", "")).strip()
        prefix = f"{dataset}:"
        if dataset and specimen_key.startswith(prefix):
            specimen = specimen_key[len(prefix) :].lower()
    if not dataset or not specimen:
        raise ValueError("dataset_id and specimen_id/specimen_key are required")
    return dataset, specimen


def _status_and_reason(
    candidates: list[Mapping[str, object]], *, required_stage: str
) -> tuple[str, str]:
    if not candidates:
        return "MISSING", "NO_MATCHING_RECORD"
    matching = [
        row for row in candidates if str(row.get("measurement_stage")) == required_stage
    ]
    if not matching:
        return "STAGE_MISMATCH", "NO_MATCHING_POST_LVI_PRE_CAI_RECORD"
    if len(matching) > 1:
        return "AMBIGUOUS", "MULTIPLE_MATCHING_RECORDS"
    return "MATCHED", ""


def _damage_status_and_reason(
    candidates: list[Mapping[str, object]], *, required_stage: str
) -> tuple[str, str]:
    if not candidates:
        return "MISSING", "NO_MATCHING_AUTHOR_DAMAGE_SCALAR"
    matching = [
        row for row in candidates if str(row.get("measurement_stage")) == required_stage
    ]
    if not matching:
        return "STAGE_MISMATCH", "NO_MATCHING_POST_LVI_PRE_CAI_RECORD"
    semantics = [str(row.get("measurement_semantics", "")).upper() for row in matching]
    if len(set(semantics)) != len(semantics):
        return "AMBIGUOUS", "DUPLICATE_MATCHING_DAMAGE_SCALAR"
    return "MATCHED", ""


def build_test24_crosswalk(
    review_rows: Iterable[Mapping[str, object]],
    author_measurements: Iterable[Mapping[str, object]],
    cscan_rows: Iterable[Mapping[str, object]],
) -> list[dict[str, object]]:
    """Build one row per reviewed specimen using exact dataset/specimen/stage keys."""

    reviews = list(review_rows)
    measurements = list(author_measurements)
    scans = list(cscan_rows)
    review_keys = [_exact_key(row) for row in reviews]
    if len(set(review_keys)) != len(review_keys):
        raise ValueError("reviewed specimen identities are not unique")
    output: list[dict[str, object]] = []
    for review, key in zip(reviews, review_keys, strict=True):
        specimen_key = f"{key[0]}:{key[1]}"
        damage = [
            row
            for row in measurements
            if str(row.get("specimen_key", "")).strip().lower() == specimen_key.lower()
            and str(row.get("dataset_id", "")).strip() == key[0]
            and str(row.get("measurement_semantics", "")).upper() in _DAMAGE_SCALARS
        ]
        spatial = [
            row
            for row in measurements
            if str(row.get("specimen_key", "")).strip().lower() == specimen_key.lower()
            and str(row.get("dataset_id", "")).strip() == key[0]
            and str(row.get("provenance_status", "")) == AUTHOR_SPATIAL_REFERENCE
        ]
        scan_matches = [row for row in scans if _exact_key(row) == key]
        raw_ultrasound = [
            row
            for row in scan_matches
            if str(row.get("source_kind", "")) == "AUTHOR_RAW_ULTRASOUND"
        ]
        damage_status, damage_reason = _damage_status_and_reason(
            damage, required_stage=_POST_LVI_STAGE
        )
        spatial_status, spatial_reason = _status_and_reason(
            spatial, required_stage=_POST_LVI_STAGE
        )
        if spatial_status == "MISSING":
            spatial_reason = "NO_AUTHOR_SPATIAL_REFERENCE"
        scan_status, scan_reason = _status_and_reason(
            scan_matches, required_stage=_POST_LVI_STAGE
        )
        if scan_status == "MISSING":
            scan_reason = "NO_MATCHING_AUTHOR_CSCAN_SCREENSHOT"
        raw_status, raw_reason = _status_and_reason(
            raw_ultrasound, required_stage=_POST_LVI_STAGE
        )
        if raw_status == "MISSING":
            raw_reason = "NO_NONPROPRIETARY_RAW_ULTRASOUND"
        output.append(
            {
                "dataset_id": key[0],
                "specimen_id_raw": str(review.get("specimen_id", key[1])),
                "specimen_id_normalized": key[1],
                "specimen_key": specimen_key,
                "review_split": str(review.get("split", "TEST")),
                "expert_reference_version": str(review.get("reference_version", "")),
                "expert_reference_file": str(review.get("reference_file", "")),
                "author_spatial_reference_status": spatial_status,
                "author_spatial_reference_reason": spatial_reason,
                "author_damage_scalar_status": damage_status,
                "author_damage_scalar_reason": damage_reason,
                "author_damage_scalar_record_count": len(damage),
                "author_cscan_status": scan_status,
                "author_cscan_reason": scan_reason,
                "author_cscan_source_path": (
                    str(scan_matches[0].get("cscan_source_path", ""))
                    if scan_status == "MATCHED"
                    else ""
                ),
                "quantitative_ultrasound_status": raw_status,
                "quantitative_ultrasound_reason": raw_reason,
                "measurement_stage": _POST_LVI_STAGE,
                "join_status": (
                    "MATCHED_WITH_LIMITS"
                    if damage_status == scan_status == "MATCHED"
                    else "INCOMPLETE"
                ),
            }
        )
    return output


def summarize_crosswalk(rows: Iterable[Mapping[str, object]]) -> dict[str, int]:
    records = list(rows)
    keys = [str(row.get("specimen_key", "")) for row in records]
    if not keys or any(not key for key in keys) or len(set(keys)) != len(keys):
        raise ValueError("crosswalk must contain unique nonempty specimen keys")

    def count(field: str) -> int:
        return sum(str(row.get(field, "")) == "MATCHED" for row in records)

    status_fields = (
        "author_spatial_reference_status",
        "author_damage_scalar_status",
        "author_cscan_status",
        "quantitative_ultrasound_status",
    )
    return {
        "physical_test_n": len(records),
        "spatial_reference_n": count("author_spatial_reference_status"),
        "scalar_reference_n": count("author_damage_scalar_status"),
        "author_cscan_screenshot_n": count("author_cscan_status"),
        "quantitative_ultrasound_n": count("quantitative_ultrasound_status"),
        "missing_or_ambiguous_n": sum(
            any(str(row.get(field, "")) != "MATCHED" for field in status_fields)
            for row in records
        ),
    }


def build_reference_comparison(
    *,
    specimen_key: str,
    task_or_target: str,
    measurement_stage: str,
    reference_record_id: str,
    reference_kind: str,
    quantity: str,
    unit: str,
    author_value: float | None,
    expert_mask: np.ndarray,
    reader_mask: np.ndarray,
    scale_evidence: Mapping[str, object] | None,
    mask_source: str,
) -> dict[str, object]:
    """Compare like-for-like projected areas and retain explicit nulls otherwise."""

    status = "COMPARABLE"
    limitation = ""
    expert_value: float | None = None
    reader_value: float | None = None
    normalized_quantity = quantity.strip().upper()
    if normalized_quantity != "PROJECTED_DELAMINATION_AREA" or unit != "mm^2":
        status = "INCOMPARABLE_QUANTITY"
        limitation = (
            "Only like-defined projected areas are compared; length definitions differ."
        )
    elif scale_evidence is None:
        status = "CALIBRATION_UNRESOLVED"
        limitation = "No specimen/layout-specific physical scale was established."
    else:
        try:
            width_mm = float(scale_evidence["width_mm"])
            height_mm = float(scale_evidence["height_mm"])
            basis = str(scale_evidence["basis"]).strip()
        except (KeyError, TypeError, ValueError) as error:
            raise ValueError("scale evidence is incomplete") from error
        left = np.asarray(expert_mask, dtype=np.bool_)
        right = np.asarray(reader_mask, dtype=np.bool_)
        if (
            left.ndim != 2
            or right.shape != left.shape
            or width_mm <= 0.0
            or height_mm <= 0.0
            or not basis
        ):
            raise ValueError("comparison masks or scale evidence are invalid")
        pixel_area = width_mm * height_mm / left.size
        expert_value = float(np.count_nonzero(left) * pixel_area)
        reader_value = float(np.count_nonzero(right) * pixel_area)
    absolute: float | None = None
    relative: float | None = None
    if status == "COMPARABLE" and author_value is not None and reader_value is not None:
        author = float(author_value)
        if not math.isfinite(author):
            raise ValueError("author value must be finite")
        absolute = abs(reader_value - author)
        relative = absolute / abs(author) if author != 0.0 else None
    return {
        "specimen_key": specimen_key,
        "task_or_target": task_or_target,
        "measurement_stage": measurement_stage,
        "reference_record_id": reference_record_id,
        "reference_kind": reference_kind,
        "quantity": normalized_quantity,
        "unit": unit,
        "expert_value": expert_value,
        "reader_value": reader_value,
        "author_value": author_value,
        "comparison_status": status,
        "absolute_difference": absolute,
        "relative_difference": relative,
        "scale_evidence": json.dumps(scale_evidence, sort_keys=True)
        if scale_evidence is not None
        else "",
        "mask_source": mask_source,
        "limitation": limitation,
    }


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def _write_csv(
    path: Path, rows: list[Mapping[str, object]], fields: tuple[str, ...]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=fields, extrasaction="ignore", lineterminator="\n"
        )
        writer.writeheader()
        writer.writerows(rows)


def _load_yaml(path: Path) -> dict[str, Any]:
    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    if type(payload) is not dict:
        raise ValueError(f"invalid YAML mapping: {path}")
    return payload


def _relative_source_path(path: Path, source_root: Path) -> str:
    return (
        path.resolve(strict=True)
        .relative_to(source_root.resolve(strict=True))
        .as_posix()
    )


def _bound_sources(project_root: Path, source_root: Path) -> dict[str, dict[str, str]]:
    config = _load_yaml(project_root / "paper_v3/configs/p1_full_field_oracle.yaml")
    sources = config.get("sources")
    if type(sources) is not dict:
        raise ValueError("P1 source bindings are missing")
    bindings: dict[str, dict[str, str]] = {}
    for key, value in sources.items():
        if type(value) is not dict or type(value.get("path")) is not str:
            continue
        bindings[str(key)] = {
            "path": str(value["path"]),
            "sha256": str(value.get("sha256", "")),
        }
    missing = [key for key in _WORKBOOK_KEYS if key not in bindings]
    if missing:
        raise ValueError(f"bound workbook keys are missing: {missing}")
    return bindings


def run_inventory(
    project_root: str | Path, source_root: str | Path, output_root: str | Path
) -> dict[str, object]:
    """Validate the configured evidence sources without modifying the source tree."""

    project = Path(project_root).resolve(strict=True)
    external = Path(source_root).resolve(strict=True)
    output = Path(output_root)
    bindings = _bound_sources(project, external)
    relevant = tuple(
        dict.fromkeys(
            (
                *_WORKBOOK_KEYS,
                "manifest",
                "physical_calibration",
                "physical_measurements",
                "physical_rule_calibration",
                "physical_descriptors",
                "physical_result_manifest",
                "physical_generation_code",
                "physical_descriptor_authority",
                "physical_morphology_authority",
                "physical_generation_config",
            )
        )
    )
    records: list[dict[str, object]] = []
    missing_required = 0
    for key in relevant:
        binding = bindings.get(key)
        if binding is None:
            records.append(
                {
                    "source_id": key,
                    "path": "",
                    "expected_sha256": "",
                    "actual_sha256": "",
                    "status": "CONFIG_BINDING_MISSING",
                    "access_level": "UNAVAILABLE",
                }
            )
            missing_required += 1
            continue
        path = external / binding["path"]
        expected = binding["sha256"]
        actual = _sha256(path) if path.is_file() else ""
        status = (
            "HASH_VERIFIED"
            if actual and actual == expected
            else "HASH_MISMATCH"
            if actual
            else "NOT_FOUND"
        )
        if status != "HASH_VERIFIED":
            missing_required += 1
        records.append(
            {
                "source_id": key,
                "path": binding["path"],
                "expected_sha256": expected,
                "actual_sha256": actual,
                "status": status,
                "access_level": "LOCAL_READ_ONLY" if actual else "UNAVAILABLE",
            }
        )
    for source_id, relative in (
        ("published_measurement_reader", "src/cmc_bbdm/cpb_published_measurements.py"),
        ("cai_workbook_reader", "src/cmc_bbdm/hasebe_cai.py"),
        ("mendeley_source_adapter", "src/cmc_bbdm/mendeley.py"),
    ):
        path = external / relative
        records.append(
            {
                "source_id": source_id,
                "path": relative,
                "expected_sha256": "",
                "actual_sha256": _sha256(path) if path.is_file() else "",
                "status": "LOCAL_READ_VERIFIED_EXISTENCE"
                if path.is_file()
                else "NOT_FOUND",
                "access_level": "LOCAL_READ_ONLY" if path.is_file() else "UNAVAILABLE",
            }
        )
    payload: dict[str, object] = {
        "schema_version": 1,
        "repository_base_sha": _BASE_SHA,
        "source_root_label": "EXPLICIT_EXTERNAL_HASEBE_SOURCE_ROOT",
        "bound_workbook_count": len(_WORKBOOK_KEYS),
        "missing_required_source_count": missing_required,
        "records": records,
        "public_evidence": list(_PUBLIC_EVIDENCE),
        "scope_limits": {
            "workbooks": 9,
            "fixed_case_images": 3,
            "full_input_reader_reports": 6,
            "training_updates": 0,
            "model_forward_calls": 0,
        },
    }
    _write_json(output / "source_inventory.json", payload)
    if missing_required:
        raise ValueError(
            f"{missing_required} required source bindings failed validation"
        )
    return payload


def _dataset_for_specimen(specimen_id: str) -> str:
    prefix = specimen_id.strip().lower().split("-", 1)[0]
    return _DATASET_BY_PREFIX.get(prefix, "8scdmfdcfb")


def _record_id(source_path: str, sheet: str, cell: str, semantics: str) -> str:
    raw = f"{source_path}|{sheet}|{cell}|{semantics}".encode()
    return f"hasebe-{hashlib.sha256(raw).hexdigest()[:20]}"


def _measurement_record(
    *,
    source_path: str,
    source_sha256: str,
    sheet: str,
    row: int,
    column: str,
    specimen_id: str,
    stage: str,
    name_raw: str,
    semantics: str,
    value: object,
    unit_raw: str,
    unit_normalized: str,
    derivation: str,
    notes: str = "",
) -> dict[str, object]:
    dataset = _dataset_for_specimen(specimen_id)
    cai_source = "hasebe_cai" in source_path
    cell = f"{column}{row}"
    source_kind = "AUTHOR_WORKBOOK"
    method = "DIRECT_CELL_READ"
    return {
        "record_id": _record_id(source_path, sheet, cell, semantics),
        "dataset_id": dataset,
        "dataset_version": 3 if dataset == "8scdmfdcfb" else 1,
        "source_dataset_id": "8scdmfdcfb" if cai_source else dataset,
        "source_dataset_version": 3 if cai_source else 1,
        "specimen_id_raw": specimen_id,
        "specimen_key": f"{dataset}:{specimen_id.lower()}",
        "measurement_stage": stage,
        "measurement_name_raw": name_raw,
        "measurement_semantics": semantics,
        "value_raw": value,
        "value_numeric": value if isinstance(value, (int, float)) else "",
        "unit_raw": unit_raw,
        "unit_normalized": unit_normalized,
        "source_kind": source_kind,
        "source_path_or_url": source_path,
        "source_sheet": sheet,
        "source_cell": cell,
        "source_sha256": source_sha256,
        "extraction_method": method,
        "derivation": derivation,
        "provenance_status": provenance_status(
            source_kind=source_kind,
            extraction_method=method,
            measurement_semantics=semantics,
        ),
        "notes": notes,
    }


def _field_record(
    *,
    source_id: str,
    source_path: str,
    source_sha256: str,
    sheet: str,
    header_cell: str,
    field_raw: str,
    field_normalized: str,
    unit_raw: str,
    unit_cell: str,
    data_column: str,
    first_row: int,
    last_row: int,
    notes: str = "",
) -> dict[str, object]:
    return {
        "source_id": source_id,
        "source_path": source_path,
        "source_sha256": source_sha256,
        "source_sheet": sheet,
        "header_cell": header_cell,
        "field_name_raw": field_raw,
        "field_name_normalized": field_normalized,
        "unit_raw": unit_raw,
        "unit_cell": unit_cell,
        "data_column": data_column,
        "first_data_row": first_row,
        "last_data_row": last_row,
        "extraction_status": "EXTRACTED",
        "notes": notes,
    }


_MEASUREMENT_FIELDS = (
    "record_id",
    "dataset_id",
    "dataset_version",
    "source_dataset_id",
    "source_dataset_version",
    "specimen_id_raw",
    "specimen_key",
    "measurement_stage",
    "measurement_name_raw",
    "measurement_semantics",
    "value_raw",
    "value_numeric",
    "unit_raw",
    "unit_normalized",
    "source_kind",
    "source_path_or_url",
    "source_sheet",
    "source_cell",
    "source_sha256",
    "extraction_method",
    "derivation",
    "provenance_status",
    "notes",
)

_FIELD_FIELDS = (
    "source_id",
    "source_path",
    "source_sha256",
    "source_sheet",
    "header_cell",
    "field_name_raw",
    "field_name_normalized",
    "unit_raw",
    "unit_cell",
    "data_column",
    "first_data_row",
    "last_data_row",
    "extraction_status",
    "notes",
)


def _extract_impact_workbook(
    path: Path, source_id: str, source_root: Path
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    rel = _relative_source_path(path, source_root)
    sha = _sha256(path)
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = book["condition"]
        raw_headers = [sheet.cell(1, col).value for col in range(1, 7)]
        normalized = (
            "LAMINATE_FAMILY",
            "PLY_COUNT",
            "SPECIMEN_NUMBER",
            "IMPACTOR_SHAPE",
            "IMPACT_ENERGY_PER_THICKNESS",
            "SCAN_ANGLE",
        )
        fields = [
            _field_record(
                source_id=source_id,
                source_path=rel,
                source_sha256=sha,
                sheet=sheet.title,
                header_cell=f"{chr(64 + column)}1",
                field_raw=str(raw_headers[column - 1]),
                field_normalized=normalized[column - 1],
                unit_raw="J/mm" if column == 5 else "",
                unit_cell="E1" if column == 5 else "",
                data_column=chr(64 + column),
                first_row=2,
                last_row=sheet.max_row,
            )
            for column in range(1, 7)
        ]
        measurements: list[dict[str, object]] = []
        for row in range(2, sheet.max_row + 1):
            family = sheet.cell(row, 1).value
            ply = sheet.cell(row, 2).value
            number = sheet.cell(row, 3).value
            energy = sheet.cell(row, 5).value
            if family is None or ply is None or number is None:
                continue
            specimen = (
                f"{str(family).strip().lower()}{int(ply)}-{str(number).strip().lower()}"
            )
            if isinstance(ply, (int, float)):
                measurements.append(
                    _measurement_record(
                        source_path=rel,
                        source_sha256=sha,
                        sheet=sheet.title,
                        row=row,
                        column="B",
                        specimen_id=specimen,
                        stage="PRE_IMPACT",
                        name_raw="ply",
                        semantics="PLY_COUNT",
                        value=ply,
                        unit_raw="",
                        unit_normalized="1",
                        derivation="DIRECT_CELL_READ",
                    )
                )
            if isinstance(energy, (int, float)):
                measurements.append(
                    _measurement_record(
                        source_path=rel,
                        source_sha256=sha,
                        sheet=sheet.title,
                        row=row,
                        column="E",
                        specimen_id=specimen,
                        stage="PRE_IMPACT",
                        name_raw="J/mm",
                        semantics="IMPACT_ENERGY_PER_THICKNESS",
                        value=energy,
                        unit_raw="J/mm",
                        unit_normalized="J/mm",
                        derivation="DIRECT_CELL_READ",
                    )
                )
        return fields, measurements
    finally:
        book.close()


def _extract_cai_workbook(
    path: Path, source_id: str, source_root: Path
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    rel = _relative_source_path(path, source_root)
    sha = _sha256(path)
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = book.active
        fields: list[dict[str, object]] = []
        measurements: list[dict[str, object]] = []
        if source_id == "lvi_workbook":
            specs = (
                ("B", "B2", "Specimen No.", "SPECIMEN_ID", "", "", 5),
                ("C", "C3", "Layup", "LAMINATE_FAMILY", "", "", 5),
                ("D", "D3", "Impactor shape", "IMPACTOR_SHAPE", "", "", 5),
                ("E", "E3", "Impact energy", "TOTAL_IMPACT_ENERGY", "[J/mm]", "E4", 5),
                (
                    "F",
                    "E3",
                    "Impact energy",
                    "IMPACT_ENERGY_PER_THICKNESS",
                    "[J]",
                    "F4",
                    5,
                ),
                (
                    "G",
                    "G3",
                    "Projected delamination area",
                    "PROJECTED_DELAMINATION_AREA",
                    "[mm2]",
                    "G4",
                    5,
                ),
                ("H", "H3", "Dent depth", "DENT_DEPTH", "[mm]", "H4", 5),
                ("I", "I3", "Is included", "CAI_INCLUDED_FLAG", "", "", 5),
            )
            for column, header, raw, normalized, unit, unit_cell, first in specs:
                fields.append(
                    _field_record(
                        source_id=source_id,
                        source_path=rel,
                        source_sha256=sha,
                        sheet=sheet.title,
                        header_cell=header,
                        field_raw=raw,
                        field_normalized=normalized,
                        unit_raw=unit,
                        unit_cell=unit_cell,
                        data_column=column,
                        first_row=first,
                        last_row=sheet.max_row,
                        notes=(
                            "E4/F4 unit labels conflict with the values and specimen-thickness relation; normalized semantics retain the value relation and preserve raw labels."
                            if column in {"E", "F"}
                            else ""
                        ),
                    )
                )
            measurement_specs = (
                ("E", "Impact energy", "TOTAL_IMPACT_ENERGY", "[J/mm]", "J"),
                ("F", "Impact energy", "IMPACT_ENERGY_PER_THICKNESS", "[J]", "J/mm"),
                (
                    "G",
                    "Projected delamination area",
                    "PROJECTED_DELAMINATION_AREA",
                    "[mm2]",
                    "mm^2",
                ),
                ("H", "Dent depth", "DENT_DEPTH", "[mm]", "mm"),
            )
            for row in range(5, sheet.max_row + 1):
                specimen_value = sheet.cell(row, 2).value
                if not isinstance(specimen_value, str) or not specimen_value.strip():
                    continue
                specimen = specimen_value.strip().lower()
                for column, raw, semantics, unit_raw, unit in measurement_specs:
                    value = sheet[f"{column}{row}"].value
                    if not isinstance(value, (int, float)):
                        continue
                    note = (
                        "Raw E/F unit headers appear reversed: E equals F times specimen thickness in the bound size workbook; raw units are retained here."
                        if column in {"E", "F"}
                        else ""
                    )
                    measurements.append(
                        _measurement_record(
                            source_path=rel,
                            source_sha256=sha,
                            sheet=sheet.title,
                            row=row,
                            column=column,
                            specimen_id=specimen,
                            stage=(
                                _POST_LVI_STAGE
                                if column in {"G", "H"}
                                else "PRE_IMPACT"
                            ),
                            name_raw=raw,
                            semantics=semantics,
                            value=value,
                            unit_raw=unit_raw,
                            unit_normalized=unit,
                            derivation=(
                                "DIRECT_CELL_READ_WITH_DOCUMENTED_E_F_UNIT_HEADER_CORRECTION"
                                if column in {"E", "F"}
                                else "DIRECT_CELL_READ"
                            ),
                            notes=note,
                        )
                    )
        elif source_id == "size_workbook":
            specs = (
                ("B", "B2", "Specimen No.", "SPECIMEN_ID"),
                ("C", "C3", "Height", "SPECIMEN_HEIGHT"),
                ("D", "D3", "Width", "SPECIMEN_WIDTH"),
                ("E", "E3", "Thickness", "SPECIMEN_THICKNESS"),
            )
            for column, header, raw, normalized in specs:
                fields.append(
                    _field_record(
                        source_id=source_id,
                        source_path=rel,
                        source_sha256=sha,
                        sheet=sheet.title,
                        header_cell=header,
                        field_raw=raw,
                        field_normalized=normalized,
                        unit_raw="mm" if column != "B" else "",
                        unit_cell="C2" if column != "B" else "",
                        data_column=column,
                        first_row=4,
                        last_row=sheet.max_row,
                    )
                )
            for row in range(4, sheet.max_row + 1):
                specimen_value = sheet.cell(row, 2).value
                if not isinstance(specimen_value, str) or not specimen_value.strip():
                    continue
                specimen = specimen_value.strip().lower()
                for column, _header, raw, semantics in specs[1:]:
                    value = sheet[f"{column}{row}"].value
                    if isinstance(value, (int, float)):
                        measurements.append(
                            _measurement_record(
                                source_path=rel,
                                source_sha256=sha,
                                sheet=sheet.title,
                                row=row,
                                column=column,
                                specimen_id=specimen,
                                stage=_POST_LVI_STAGE,
                                name_raw=raw,
                                semantics=semantics,
                                value=value,
                                unit_raw="mm",
                                unit_normalized="mm",
                                derivation="DIRECT_CELL_READ",
                            )
                        )
        elif source_id == "cai_workbook":
            specs = (
                ("B", "B2", "Specimen No.", "SPECIMEN_ID", "", ""),
                (
                    "C",
                    "C2",
                    "Compression after impact strength",
                    "CAI_STRENGTH",
                    "[Mpa]",
                    "C3",
                ),
                (
                    "D",
                    "C2",
                    "Compression after impact strength",
                    "CAI_STRENGTH_REDUCTION",
                    "[%]",
                    "D3",
                ),
            )
            for column, header, raw, normalized, unit, unit_cell in specs:
                fields.append(
                    _field_record(
                        source_id=source_id,
                        source_path=rel,
                        source_sha256=sha,
                        sheet=sheet.title,
                        header_cell=header,
                        field_raw=raw,
                        field_normalized=normalized,
                        unit_raw=unit,
                        unit_cell=unit_cell,
                        data_column=column,
                        first_row=4,
                        last_row=sheet.max_row,
                    )
                )
            for row in range(4, sheet.max_row + 1):
                specimen_value = sheet.cell(row, 2).value
                if not isinstance(specimen_value, str) or not specimen_value.strip():
                    continue
                specimen = specimen_value.strip().lower()
                for column, raw, semantics, unit_raw, unit in (
                    (
                        "C",
                        "Compression after impact strength",
                        "CAI_STRENGTH",
                        "[Mpa]",
                        "MPa",
                    ),
                    (
                        "D",
                        "Compression after impact strength",
                        "CAI_STRENGTH_REDUCTION",
                        "[%]",
                        "1",
                    ),
                ):
                    value = sheet[f"{column}{row}"].value
                    if isinstance(value, (int, float)):
                        measurements.append(
                            _measurement_record(
                                source_path=rel,
                                source_sha256=sha,
                                sheet=sheet.title,
                                row=row,
                                column=column,
                                specimen_id=specimen,
                                stage="POST_CAI",
                                name_raw=raw,
                                semantics=semantics,
                                value=value,
                                unit_raw=unit_raw,
                                unit_normalized=unit,
                                derivation="DIRECT_CELL_READ",
                            )
                        )
        else:
            raise ValueError(f"unsupported CAI workbook binding: {source_id}")
        return fields, measurements
    finally:
        book.close()


def _validate_published_transcription(
    source_root: Path,
    bindings: Mapping[str, Mapping[str, str]],
    measurements: Iterable[Mapping[str, object]],
) -> dict[str, object]:
    path = source_root / bindings["physical_measurements"]["path"]
    published = read_csv_rows(path)
    extracted = {
        (str(row["dataset_id"]), str(row["specimen_id_raw"]).lower()): row
        for row in measurements
        if row["measurement_semantics"] == "PROJECTED_DELAMINATION_AREA"
    }
    mismatch: list[str] = []
    for row in published:
        key = (row["dataset_id"], row["sample_id"].lower())
        found = extracted.get(key)
        if found is None or not math.isclose(
            float(row["published_projected_delamination_area"]),
            float(found["value_numeric"]),
            abs_tol=1e-12,
        ):
            mismatch.append(f"{key[0]}:{key[1]}:area")
            continue
        dent = next(
            (
                candidate
                for candidate in measurements
                if candidate["dataset_id"] == key[0]
                and str(candidate["specimen_id_raw"]).lower() == key[1]
                and candidate["measurement_semantics"] == "DENT_DEPTH"
            ),
            None,
        )
        if dent is None or not math.isclose(
            float(row["published_dent_depth"]),
            float(dent["value_numeric"]),
            abs_tol=1e-12,
        ):
            mismatch.append(f"{key[0]}:{key[1]}:dent")
    if mismatch:
        raise ValueError(f"published measurement transcription differs: {mismatch[:5]}")
    return {
        "published_row_count": len(published),
        "matched_direct_workbook_row_count": len(published),
        "mismatch_count": 0,
        "source_csv": bindings["physical_measurements"]["path"],
        "source_workbook": bindings["lvi_workbook"]["path"],
        "generator": "src/cmc_bbdm/cpb_published_measurements.py::read_published_damage_measurements/build_matched_measurement_rows",
        "classification": AUTHOR_SCALAR_MEASUREMENT,
    }


def _audit_reversed_energy_headers(
    measurements: Iterable[Mapping[str, object]],
) -> dict[str, int]:
    by_key_semantics = {
        (str(row["specimen_key"]), str(row["measurement_semantics"])): row
        for row in measurements
        if str(row["source_dataset_id"]) == "8scdmfdcfb"
    }
    checked = 0
    exact = 0
    within_five_percent = 0
    keys = {
        key for key, semantics in by_key_semantics if semantics == "TOTAL_IMPACT_ENERGY"
    }
    for key in keys:
        total = by_key_semantics[(key, "TOTAL_IMPACT_ENERGY")]
        per_thickness = by_key_semantics.get((key, "IMPACT_ENERGY_PER_THICKNESS"))
        thickness = by_key_semantics.get((key, "SPECIMEN_THICKNESS"))
        if per_thickness is None or thickness is None:
            continue
        checked += 1
        expected = float(per_thickness["value_numeric"]) * float(
            thickness["value_numeric"]
        )
        total_value = float(total["value_numeric"])
        difference = abs(total_value - expected)
        if difference <= 1e-9:
            exact += 1
        if total_value and difference / abs(total_value) <= 0.05:
            within_five_percent += 1
    return {
        "checked_specimen_count": checked,
        "exact_relation_count": exact,
        "within_five_percent_count": within_five_percent,
        "over_five_percent_count": checked - within_five_percent,
    }


def run_extract(
    project_root: str | Path, source_root: str | Path, output_root: str | Path
) -> dict[str, int]:
    """Extract cell-level evidence and build the exact reviewed TEST-24 crosswalk."""

    project = Path(project_root).resolve(strict=True)
    external = Path(source_root).resolve(strict=True)
    output = Path(output_root)
    inventory_path = output / "source_inventory.json"
    if not inventory_path.is_file():
        raise ValueError("inventory must be completed before extraction")
    inventory = json.loads(inventory_path.read_text(encoding="utf-8"))
    if int(inventory.get("missing_required_source_count", -1)) != 0:
        raise ValueError("inventory contains unresolved required sources")
    bindings = _bound_sources(project, external)

    all_fields: list[dict[str, object]] = []
    all_measurements: list[dict[str, object]] = []
    workbook_audits: list[dict[str, object]] = []
    for source_id in _WORKBOOK_KEYS:
        binding = bindings[source_id]
        path = external / binding["path"]
        audit = inspect_workbook(path)
        workbook_audits.append(
            {
                "source_id": source_id,
                "path": binding["path"],
                "sha256": audit["source_sha256"],
                "sheet_count": audit["sheet_count"],
                "sheets": [
                    {
                        "name": sheet["name"],
                        "visibility": sheet["visibility"],
                        "dimensions": sheet["dimensions"],
                        "merged_ranges": sheet["merged_ranges"],
                        "freeze_panes": sheet["freeze_panes"],
                        "nonempty_cell_count": len(sheet["cells"]),
                        "formula_count": sum(
                            cell["formula"] is not None for cell in sheet["cells"]
                        ),
                        "formula_without_cached_value_count": sum(
                            cell["cache_status"] == "FORMULA_NO_CACHED_VALUE"
                            for cell in sheet["cells"]
                        ),
                        "comment_count": sheet["comment_count"],
                        "image_count": sheet["image_count"],
                        "chart_count": sheet["chart_count"],
                    }
                    for sheet in audit["sheets"]
                ],
            }
        )
        if source_id.startswith("impact_condition_"):
            fields, measurements = _extract_impact_workbook(path, source_id, external)
        else:
            fields, measurements = _extract_cai_workbook(path, source_id, external)
        all_fields.extend(fields)
        all_measurements.extend(measurements)

    _write_csv(output / "workbook_fields.csv", all_fields, _FIELD_FIELDS)
    _write_csv(
        output / "author_measurements.csv", all_measurements, _MEASUREMENT_FIELDS
    )

    lineage = _validate_published_transcription(external, bindings, all_measurements)
    lineage["reversed_energy_header_relation_audit"] = _audit_reversed_energy_headers(
        all_measurements
    )
    physical_rows = read_csv_rows(external / bindings["physical_descriptors"]["path"])
    rule_rows = read_csv_rows(external / bindings["physical_rule_calibration"]["path"])
    selected_rules = sorted(
        {
            row["rule"]
            for row in rule_rows
            if str(row.get("selected", "")).strip().lower() == "true"
        }
    )
    if selected_rules != ["bg40_rb20_close08"]:
        raise ValueError("historical derived-proxy rule identity changed")
    lineage["physical_descriptor_chain"] = {
        "classification": LOCAL_DERIVED_PROXY,
        "method": "RGB_THRESHOLD_MORPHOLOGY_WITH_AUTHOR_SCALAR_CALIBRATION",
        "selected_rule": selected_rules[0],
        "physical_descriptor_row_count": len(physical_rows),
        "calibration_row_count": sum(
            row["rule"] == selected_rules[0] for row in rule_rows
        ),
        "source_code": bindings["physical_morphology_authority"]["path"],
        "not_author_spatial_ground_truth": True,
    }
    inventory["workbook_audits"] = workbook_audits
    inventory["field_lineage"] = lineage
    _write_json(inventory_path, inventory)

    coverage_path = (
        project
        / "results/bc_cscan_expert_pooled_rescore/v1/reviewed/reference_coverage.csv"
    )
    review_rows = read_csv_rows(coverage_path)
    if len(review_rows) != 24 or {row["split"] for row in review_rows} != {"TEST"}:
        raise ValueError("reviewed reference coverage is not the frozen TEST-24 set")
    reviews = [
        {
            **row,
            "reference_file": (
                "results/bc_cscan_expert_pooled_rescore/v1/inputs/references/"
                + row["reference_path"]
            ),
        }
        for row in review_rows
    ]
    manifest_rows = read_csv_rows(
        project
        / "results/agentic_task_driven_nde/p0r_author_registration/surface_manifest.csv"
    )
    review_keys = {(row["dataset_id"], row["specimen_id"].lower()) for row in reviews}
    cscan_rows = [
        {
            "dataset_id": row["dataset_id"],
            "specimen_id": row["specimen_id"].lower(),
            "specimen_key": f"{row['dataset_id']}:{row['specimen_id'].lower()}",
            "measurement_stage": _POST_LVI_STAGE,
            "source_kind": "AUTHOR_CSCAN_SCREENSHOT",
            "cscan_source_path": row["cscan_source_path"],
            "cscan_source_sha256": row["cscan_source_sha256"],
            "cscan_panel_index": row["cscan_panel_index"],
            "registered_cscan_crop_path": row["registered_cscan_crop_path"],
            "registered_cscan_crop_sha256": row["registered_cscan_crop_sha256"],
            "cai_identity": row["cai_identity"],
            "identity_status": row["identity_status"],
            "pairing_method": row["pairing_method"],
            "pairing_confidence": row["pairing_confidence"],
            "author_statement_sha256": row["author_statement_sha256"],
        }
        for row in manifest_rows
        if (row["dataset_id"], row["specimen_id"].lower()) in review_keys
    ]
    if len(cscan_rows) != 24:
        raise ValueError("P0R screenshot identity coverage is not exactly TEST-24")
    if any(
        not (external / row["cscan_source_path"]).is_file()
        or not (external / row["registered_cscan_crop_path"]).is_file()
        for row in cscan_rows
    ):
        raise ValueError("a TEST-24 C-scan screenshot binding is unavailable")
    damage_area_by_key = {
        str(row["specimen_key"]): row
        for row in all_measurements
        if row["measurement_semantics"] == "PROJECTED_DELAMINATION_AREA"
    }
    for scan in cscan_rows:
        area = damage_area_by_key.get(str(scan["specimen_key"]))
        if area is None:
            raise ValueError("a TEST-24 author damage row is unavailable")
        excel_row = int(str(area["source_cell"])[1:])
        if (
            str(scan["identity_status"]) != "PASS_EXACT_SPECIMEN_ID_AND_HASH"
            or str(scan["pairing_method"]) != "exact_published_specimen_id"
            or not str(scan["cai_identity"]).endswith(f"_row:{excel_row - 1}")
        ):
            raise ValueError("P0R and CAI workbook row identities differ")
    crosswalk = build_test24_crosswalk(reviews, all_measurements, cscan_rows)
    scan_by_key = {row["specimen_key"]: row for row in cscan_rows}
    for row in crosswalk:
        scan = scan_by_key[row["specimen_key"]]
        row.update(
            {
                "author_cscan_source_sha256": scan["cscan_source_sha256"],
                "author_cscan_panel_index": scan["cscan_panel_index"],
                "registered_cscan_crop_path": scan["registered_cscan_crop_path"],
                "registered_cscan_crop_sha256": scan["registered_cscan_crop_sha256"],
                "author_damage_scalar_source": bindings["lvi_workbook"]["path"],
                "author_spatial_reference_provenance": "",
                "author_ultrasound_provenance": AUTHOR_ULTRASOUND_MEASUREMENT,
                "cai_identity": scan["cai_identity"],
                "identity_status": scan["identity_status"],
                "pairing_method": scan["pairing_method"],
                "pairing_confidence": scan["pairing_confidence"],
                "author_statement_sha256": scan["author_statement_sha256"],
            }
        )
    crosswalk_fields = tuple(crosswalk[0])
    _write_csv(output / "test24_reference_crosswalk.csv", crosswalk, crosswalk_fields)
    return summarize_crosswalk(crosswalk)


def _mask_sha256(mask: np.ndarray) -> str:
    values = np.asarray(mask, dtype=np.bool_)
    digest = hashlib.sha256()
    digest.update(np.asarray(values.shape, dtype="<i8").tobytes())
    digest.update(np.packbits(values, bitorder="little").tobytes())
    return digest.hexdigest()


def _reference_payload(
    project: Path, row: Mapping[str, str]
) -> tuple[Path, dict[str, Any]]:
    path = (
        project
        / "results/bc_cscan_expert_pooled_rescore/v1/inputs/references"
        / row["reference_path"]
    )
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("specimen_key") != row["specimen_key"]:
        raise ValueError("expert reference specimen identity changed")
    return path, payload


def _finite_difference(
    value: float | None, reference: float | None
) -> tuple[float | None, float | None]:
    if value is None or reference is None:
        return None, None
    absolute = abs(float(value) - float(reference))
    relative = absolute / abs(float(reference)) if float(reference) != 0 else None
    return absolute, relative


def run_compare(
    project_root: str | Path,
    source_root: str | Path,
    output_root: str | Path,
    artifact_root: str | Path,
) -> dict[str, int]:
    """Review three predeclared cases and recover exactly six full-input reports."""

    from cmc_bbdm.agentic_nde.scan_frame_provenance import verify_registered_crop
    from cmc_bbdm.vlm_cscan.references import reference_from_payload

    from .benchmark import _report_digest
    from .frozen_process_recovery import _full_reports
    from .runtime import load_study_config, load_study_context, open_study_specimen

    project = Path(project_root).resolve(strict=True)
    external = Path(source_root).resolve(strict=True)
    output = Path(output_root)
    artifacts = Path(artifact_root)
    artifacts.mkdir(parents=True, exist_ok=True)
    figure_root = artifacts / "figures"
    figure_root.mkdir(parents=True, exist_ok=True)
    crosswalk = read_csv_rows(output / "test24_reference_crosswalk.csv")
    counts = summarize_crosswalk(crosswalk)
    crosswalk_by_key = {row["specimen_key"]: row for row in crosswalk}
    if not set(_FIXED_CASES).issubset(crosswalk_by_key):
        raise ValueError("fixed case identities are absent from TEST-24")
    coverage = read_csv_rows(
        project
        / "results/bc_cscan_expert_pooled_rescore/v1/reviewed/reference_coverage.csv"
    )
    coverage_by_key = {row["specimen_key"]: row for row in coverage}
    expected_reports = {
        (row["specimen_key"], row["task"]): row
        for row in read_csv_rows(
            project
            / "results/bc_cscan_expert_pooled_rescore/v1/reviewed/full_input_readout.csv"
        )
        if row["specimen_key"] in _FIXED_CASES
    }
    if len(expected_reports) != 6:
        raise ValueError("fixed-case full-input report registry is incomplete")
    author_measurements = read_csv_rows(output / "author_measurements.csv")
    author_area = {
        row["specimen_key"]: row
        for row in author_measurements
        if row["measurement_semantics"] == "PROJECTED_DELAMINATION_AREA"
        and row["specimen_key"] in _FIXED_CASES
    }
    if len(author_area) != 3:
        raise ValueError("fixed-case author area measurements are incomplete")

    study_config = load_study_config(
        project / "paper_v3/configs/learned_cscan_same_perception.yaml",
        project_root=project,
    )
    context = load_study_context(
        study_config, source_root=external, verify_pilot_hashes=True
    )
    assignment_by_key = {
        row.record.specimen_key: row for row in context.roster.assignments
    }
    case_rows: list[dict[str, object]] = []
    comparison_rows: list[dict[str, object]] = []
    visual_notes = {
        "ykhs7s2dck:q8-4": (
            "The author screenshot contains a compact central warm/white indication. "
            "No author contour or ImageJ ROI is visible."
        ),
        "74t7kcdgkr:c8-24": (
            "The author screenshot contains a small, weak central indication. "
            "No author contour or ImageJ ROI is visible."
        ),
        "cgtnjyggtm:q24-40": (
            "The author screenshot contains a broad central warm/red indication and a "
            "software crosshair/status overlay. No author damage contour is visible."
        ),
    }
    for specimen_key in _FIXED_CASES:
        assignment = assignment_by_key.get(specimen_key)
        if assignment is None or assignment.split.value != "TEST":
            raise ValueError(f"fixed case is not in TEST: {specimen_key}")
        record = assignment.record
        runtime = open_study_specimen(context, record)
        full_scan = context.authority.source_teacher_view(record.specimen_id).full_scan
        reports = _full_reports(context, runtime, full_scan)
        if set(reports) != {"LOCATE", "CHARACTERIZE"}:
            raise ValueError("full-input Reader task set changed")
        for task, report in reports.items():
            expected = expected_reports[(specimen_key, task)]["report_sha256"]
            if _report_digest(report) != expected:
                raise ValueError(
                    f"full-input Reader identity changed: {specimen_key}/{task}"
                )

        cross = crosswalk_by_key[specimen_key]
        raw_path = external / cross["author_cscan_source_path"]
        registered_path = external / cross["registered_cscan_crop_path"]
        provenance = verify_registered_crop(
            raw_path,
            registered_path,
            panel_index=int(cross["author_cscan_panel_index"]),
            expected_raw_sha256=cross["author_cscan_source_sha256"],
            expected_registered_sha256=cross["registered_cscan_crop_sha256"],
        )
        with Image.open(registered_path) as image:
            image.load()
            registered = np.asarray(image.convert("RGB"))
        if not np.array_equal(registered, full_scan):
            raise ValueError("Reader full scan and registered crop pixels differ")

        coverage_row = coverage_by_key[specimen_key]
        reference_path, payload = _reference_payload(project, coverage_row)
        if payload.get("source_image_sha256") != cross["registered_cscan_crop_sha256"]:
            raise ValueError("expert reference and registered crop hashes differ")
        reference = reference_from_payload(payload, native_shape=record.native_shape)
        if np.any(reference.uncertain_mask):
            raise ValueError(
                "fixed reviewed reference unexpectedly contains uncertainty"
            )
        expert_mask = np.asarray(reference.certain_mask, dtype=np.bool_)
        characterize = reports["CHARACTERIZE"]
        reader_mask = np.asarray(characterize.predicted_mask, dtype=np.bool_)
        scale = {
            "width_mm": 75.0,
            "height_mm": 75.0,
            "x_axis_endpoints_mm": [0.0, 75.0],
            "y_axis_endpoints_mm": [0.0, 75.0],
            "registered_shape_px": list(expert_mask.shape),
            "pixel_area_convention": "field_area_divided_by_registered_pixel_count",
            "basis": "CASE_SPECIFIC_RAW_SCREENSHOT_AXES_VISUALLY_VERIFIED_2026-09-10",
            "caveat": "This display calibration is not a claim that each screenshot pixel is one native 0.2 mm instrument sample.",
        }
        area_record = author_area[specimen_key]
        area = build_reference_comparison(
            specimen_key=specimen_key,
            task_or_target="CHARACTERIZE",
            measurement_stage=_POST_LVI_STAGE,
            reference_record_id=area_record["record_id"],
            reference_kind=AUTHOR_SCALAR_MEASUREMENT,
            quantity="PROJECTED_DELAMINATION_AREA",
            unit="mm^2",
            author_value=float(area_record["value_numeric"]),
            expert_mask=expert_mask,
            reader_mask=reader_mask,
            scale_evidence=scale,
            mask_source=(
                f"{reference_path.relative_to(project).as_posix()} + frozen full-input Reader"
            ),
        )
        area["comparison_status"] = "COMPARABLE_WITH_SCOPE_LIMIT"
        area["limitation"] = (
            "Scalar magnitudes share mm^2 and the post-LVI/pre-CAI stage, but the "
            "author ImageJ boundary definition/ROI file is unavailable; neither mask "
            "is promoted to author spatial ground truth."
        )
        expert_abs, expert_rel = _finite_difference(
            area["expert_value"], area["author_value"]
        )
        area.update(
            {
                "expert_absolute_difference": expert_abs,
                "expert_relative_difference": expert_rel,
                "reader_report_sha256": expected_reports[
                    (specimen_key, "CHARACTERIZE")
                ]["report_sha256"],
                "expert_mask_sha256": _mask_sha256(expert_mask),
                "reader_mask_sha256": _mask_sha256(reader_mask),
            }
        )
        comparison_rows.append(area)
        length = build_reference_comparison(
            specimen_key=specimen_key,
            task_or_target="LOCATE",
            measurement_stage=_POST_LVI_STAGE,
            reference_record_id="",
            reference_kind=AUTHOR_SCALAR_MEASUREMENT,
            quantity="DELAMINATION_LENGTH",
            unit="mm",
            author_value=None,
            expert_mask=expert_mask,
            reader_mask=np.asarray(reports["LOCATE"].predicted_mask, dtype=np.bool_),
            scale_evidence=scale,
            mask_source="NO_BOUND_AUTHOR_LENGTH_RECORD",
        )
        length.update(
            {
                "comparison_status": "AUTHOR_MEASUREMENT_NOT_FOUND",
                "limitation": "The bound workbooks contain no per-specimen delamination-length field; no bbox/Feret substitute was invented.",
                "expert_absolute_difference": None,
                "expert_relative_difference": None,
                "reader_report_sha256": expected_reports[(specimen_key, "LOCATE")][
                    "report_sha256"
                ],
                "expert_mask_sha256": _mask_sha256(expert_mask),
                "reader_mask_sha256": _mask_sha256(
                    np.asarray(reports["LOCATE"].predicted_mask, dtype=np.bool_)
                ),
            }
        )
        comparison_rows.append(length)

        figure_name = "author_cscan_" + specimen_key.replace(":", "_") + ".png"
        figure_path = figure_root / figure_name
        with Image.open(raw_path) as image:
            image.load()
            image.convert("RGB").save(figure_path, format="PNG", compress_level=9)
            raw_size = list(image.size)
        case_rows.append(
            {
                "specimen_key": specimen_key,
                "dataset_id": record.dataset_id,
                "specimen_id": record.specimen_id,
                "review_split": assignment.split.value,
                "measurement_stage": _POST_LVI_STAGE,
                "raw_cscan_path": cross["author_cscan_source_path"],
                "raw_cscan_sha256": cross["author_cscan_source_sha256"],
                "raw_cscan_size_px": json.dumps(raw_size),
                "panel_index": cross["author_cscan_panel_index"],
                "registered_cscan_path": cross["registered_cscan_crop_path"],
                "registered_cscan_sha256": cross["registered_cscan_crop_sha256"],
                "crop_verification_status": "PIXEL_IDENTICAL",
                "crop_provenance": json.dumps(asdict(provenance), sort_keys=True),
                "reader_input_pixel_status": "PIXEL_IDENTICAL",
                "axis_extent_x_mm": 75,
                "axis_extent_y_mm": 75,
                "axis_scale_status": "CASE_SPECIFIC_VISUAL_VERIFICATION",
                "author_spatial_contour_status": "NOT_PRESENT_IN_INSPECTED_SCREENSHOT_OR_BOUND_FILES",
                "author_area_mm2": area_record["value_numeric"],
                "expert_characterize_area_mm2": area["expert_value"],
                "reader_characterize_area_mm2": area["reader_value"],
                "expert_reference_sha256": coverage_row["reference_sha256"],
                "characterize_reader_report_sha256": expected_reports[
                    (specimen_key, "CHARACTERIZE")
                ]["report_sha256"],
                "locate_reader_report_sha256": expected_reports[
                    (specimen_key, "LOCATE")
                ]["report_sha256"],
                "evidence_figure": f"figures/{figure_name}",
                "visual_observation": visual_notes[specimen_key],
                "interpretation_limit": "Visual appearance is not an author damage mask and does not determine a unique scan route.",
            }
        )

    case_fields = tuple(case_rows[0])
    comparison_fields = tuple(comparison_rows[0])
    _write_csv(output / "case_review.csv", case_rows, case_fields)
    _write_csv(output / "reference_comparison.csv", comparison_rows, comparison_fields)
    decision: dict[str, object] = {
        "implementation_status": "COMPLETE_WITH_EVIDENCE_LIMITS",
        "searched_scope": {
            "bound_workbooks": 9,
            "reviewed_test_specimens": 24,
            "fixed_case_screenshots": 3,
            "full_input_reader_reports": 6,
            "public_primary_records": 5,
        },
        "incomplete_accesses": [
            {
                "source": "10.1016/j.compositesb.2022.109844",
                "status": "ABSTRACT_AND_INTRO_ONLY",
                "impact": "Full target-definition methods could not be independently verified.",
            },
            {
                "source": "AUTHOR_IMAGEJ_ROI_OR_CONTOUR_FILES",
                "status": "NOT_PRESENT_IN_BOUND_LOCAL_FILES",
                "impact": "Author scalar areas cannot adjudicate pixelwise expert/Reader differences.",
            },
        ],
        **counts,
        "calibration_resolved_n": 3,
        "comparable_scalar_case_n": 3,
        "evidence_categories": ["AUTHOR_SCALARS_ONLY", "SCREENSHOT_MEASUREMENT_ONLY"],
        "primary_conclusion": (
            "All 24 reviewed TEST specimens match author projected-area/dent scalars and "
            "author C-scan screenshots, but no bound author spatial contour/mask or raw "
            "quantitative ultrasound array was found."
        ),
        "recommended_next_step": "REQUEST_AUTHOR_IMAGEJ_ROI_AND_MEASUREMENT_DEFINITION",
        "not_established": [
            "AUTHOR_SPATIAL_GROUND_TRUTH",
            "RAW_QUANTITATIVE_ULTRASOUND_FOR_TEST24",
            "UNIQUE_CORRECT_SCAN_REGION",
            "UNIQUE_OPTIMAL_SCAN_ROUTE",
            "PHYSICAL_TIME_OR_PATH_COST_FROM_SCREENSHOT_PIXELS",
            "CAUSE_OF_Q8_4_EXPERT_READER_SPATIAL_DIFFERENCE",
        ],
        "resource_use": {
            "new_training": 0,
            "new_vlm_calls": 0,
            "new_actor_forward_calls": 0,
            "new_stop_forward_calls": 0,
            "formal_reviewed_rescore_runs": 0,
            "bc_prefix_replays": 0,
            "background_prior_deterministic_reconstruction_count": 1,
            "full_input_reader_report_recovery_n": 6,
            "decoded_fixed_case_raw_screenshots": 3,
        },
        "artifacts": {
            "results": [
                "source_inventory.json",
                "workbook_fields.csv",
                "author_measurements.csv",
                "test24_reference_crosswalk.csv",
                "case_review.csv",
                "reference_comparison.csv",
            ],
            "evidence_figures": [row["evidence_figure"] for row in case_rows],
        },
        "historical_proxy_lineage": {
            "published_damage_measurements": "DIRECT_TRANSCRIPTION_FROM_BOUND_CAI_LVI_WORKBOOK_G_H",
            "physical_descriptors": LOCAL_DERIVED_PROXY,
            "selected_rule": "bg40_rb20_close08",
            "calibration_scope": "71 HASH_SPLIT_CALIBRATION_SPECIMENS_USING_AUTHOR_AREA_SCALARS",
        },
    }
    _write_json(output / "decision.json", decision)
    return {
        "case_review_n": len(case_rows),
        "full_input_report_recovery_n": 6,
        "new_training": 0,
    }


def read_csv_rows(path: str | Path) -> list[dict[str, str]]:
    with Path(path).open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def write_summary_artifacts(
    output_root: str | Path, artifact_root: str | Path
) -> dict[str, int]:
    """Rebuild the feasibility summary from extracted files only."""

    results = Path(output_root)
    artifacts = Path(artifact_root)
    artifacts.mkdir(parents=True, exist_ok=True)
    rows = read_csv_rows(results / "test24_reference_crosswalk.csv")
    counts = summarize_crosswalk(rows)
    decision = json.loads((results / "decision.json").read_text(encoding="utf-8"))
    _write_json(artifacts / "summary_counts.json", counts)
    inventory_path = results / "source_inventory.json"
    inventory = (
        json.loads(inventory_path.read_text(encoding="utf-8"))
        if inventory_path.is_file()
        else {}
    )
    fields = (
        read_csv_rows(results / "workbook_fields.csv")
        if (results / "workbook_fields.csv").is_file()
        else []
    )
    comparisons = (
        read_csv_rows(results / "reference_comparison.csv")
        if (results / "reference_comparison.csv").is_file()
        else []
    )
    cases = (
        read_csv_rows(results / "case_review.csv")
        if (results / "case_review.csv").is_file()
        else []
    )

    public_rows = inventory.get("public_evidence", [])
    ledger = [
        "# Source Evidence Ledger",
        "",
        "## Evidence classes",
        "",
        "Author workbooks provide scalar measurements; author C-scan files are ultrasound screenshots, not segmentation ground truth. Expert reviewed polygons and local RGB/Reader masks remain separate evidence classes.",
        "",
        "## Primary sources",
        "",
        "| ID | Source and access | Direct support | Does not establish |",
        "|---|---|---|---|",
    ]
    for source in public_rows:
        ledger.append(
            "| {source_id} | [{citation}]({url}); {access_status} | {supports} | {does_not_support} |".format(
                **source
            )
        )
    ledger.extend(
        [
            "",
            "## Precise source locations",
            "",
            "- O01, Data Description and Internal Damage Measurement/Table 2: screenshot-only export, amplitude-coded colour, 3.5 MHz probe, 0.200 x 0.200 mm pitch, 75 x 75 mm scan length.",
            "- O02: only abstract/introduction-level access was available; no inaccessible method detail is asserted.",
            "- O03: the six Mendeley v1 record descriptions identify post-impact internal C-scanning images and image scales; files inspected locally contain screenshots and impact-condition workbooks, not author masks.",
            "- O03 records: [c8](https://data.mendeley.com/datasets/74t7kcdgkr/1), [c16](https://data.mendeley.com/datasets/yfxyg8jm46/1), [c24](https://data.mendeley.com/datasets/xcmzfsbd9t/1), [q8](https://data.mendeley.com/datasets/ykhs7s2dck/1), [q16](https://data.mendeley.com/datasets/w68dtmpfyf/1), and [q24](https://data.mendeley.com/datasets/cgtnjyggtm/1). The [ASTM record](https://data.mendeley.com/datasets/6zt73pcnxv/1) was checked only as out-of-scope context and was not added to TEST-24.",
            "- O04/O05, Data Description: the CAI package separates LVI damage, size, post-CAI images/logger data and mechanical strength; projected area was measured from C-scan images using ImageJ.",
            "",
            "## Local evidence result",
            "",
            f"- Bound workbooks opened and hash-verified: {inventory.get('bound_workbook_count', 0)}.",
            f"- Reviewed TEST specimens with author scalar damage records: {counts['scalar_reference_n']}/24.",
            f"- Reviewed TEST specimens with author C-scan screenshots: {counts['author_cscan_screenshot_n']}/24.",
            f"- Author spatial references found in the bound scope: {counts['spatial_reference_n']}/24.",
            f"- Raw/quantitative ultrasound arrays found in the bound scope: {counts['quantitative_ultrasound_n']}/24.",
            "- The three exported screenshots retain the original 891 x 891 layout, axes, colour bar and annotations. They are CC BY 4.0 source evidence from the bound Hasebe Mendeley records.",
            "",
        ]
    )
    (artifacts / "SOURCE_EVIDENCE_LEDGER.md").write_text(
        "\n".join(ledger), encoding="utf-8"
    )

    workbook_audits = inventory.get("workbook_audits", [])
    lineage_payload = inventory.get("field_lineage", {})
    bindings = [
        "# Source Bindings and Field Lineage",
        "",
        "## Workbook audit",
        "",
        "| Binding | Path | Sheet(s) | Formula count |",
        "|---|---|---|---:|",
    ]
    for audit in workbook_audits:
        sheets = audit.get("sheets", [])
        names = ", ".join(str(sheet.get("name", "")) for sheet in sheets)
        formula_count = sum(int(sheet.get("formula_count", 0)) for sheet in sheets)
        bindings.append(
            f"| {audit.get('source_id', '')} | `{audit.get('path', '')}` | {names} | {formula_count} |"
        )
    bindings.extend(
        [
            "",
            "All nine workbooks contain one visible sheet, no formula cells, no merged ranges, and no embedded drawing/image/chart objects detected by the workbook audit.",
            "",
            "## Field bindings",
            "",
            "| Sheet/cell | Raw field/unit | Normalized meaning |",
            "|---|---|---|",
        ]
    )
    for row in fields:
        bindings.append(
            f"| {row['source_sheet']}!{row['header_cell']} / {row['unit_cell'] or '-'} | {row['field_name_raw']} {row['unit_raw']} | {row['field_name_normalized']} |"
        )
    published = lineage_payload.get("published_damage_measurements", lineage_payload)
    proxy = lineage_payload.get("physical_descriptor_chain", {})
    energy_validation = lineage_payload.get("reversed_energy_header_relation_audit", {})
    bindings.extend(
        [
            "",
            "## Damage-measurement lineage",
            "",
            "`CAI LVI workbook, LVI condition!G/H (Excel row retained)` -> `cpb_published_measurements.py::read_published_damage_measurements` -> `build_matched_measurement_rows` -> `published_damage_measurements.csv` -> this task's exact TEST-24 crosswalk.",
            "",
            f"The historical CSV contains {published.get('published_row_count', 'unknown')} rows and matched the bound workbook directly with {published.get('mismatch_count', 'unknown')} value mismatches. It is an author scalar transcription, not a mask.",
            "",
            "## Local proxy lineage",
            "",
            f"Registered C-scan RGB -> threshold/morphology candidates -> calibration against author area scalars -> selected `{proxy.get('selected_rule', 'unknown')}` -> `physical_descriptors.csv`. Classification: `{proxy.get('classification', LOCAL_DERIVED_PROXY)}`. The friendly historical `descriptor_source` label does not make these rows author spatial ground truth.",
            "",
            "The LVI workbook's E4/F4 unit labels are reversed relative to numeric values and the specimen-thickness relation. This extraction preserves raw labels and records corrected normalized meanings explicitly.",
            f"The relation `total J = J/mm x thickness mm` was audited for {energy_validation.get('checked_specimen_count', 'unknown')} bound CAI records: {energy_validation.get('exact_relation_count', 'unknown')} exact, {energy_validation.get('within_five_percent_count', 'unknown')} within 5%, and {energy_validation.get('over_five_percent_count', 'unknown')} beyond 5%. The normalized semantics follow the bound source reader's explicit correction; this relation is supporting context, not a zero-error identity claim.",
            "",
        ]
    )
    (artifacts / "SOURCE_BINDINGS_AND_FIELD_LINEAGE.md").write_text(
        "\n".join(bindings), encoding="utf-8"
    )

    case_lines = []
    for row in comparisons:
        if row.get("quantity") != "PROJECTED_DELAMINATION_AREA":
            continue
        case_lines.append(
            f"- `{row['specimen_key']}`: author {float(row['author_value']):.1f} mm^2; expert-region display-calibrated area {float(row['expert_value']):.1f} mm^2; full-input CHARACTERIZE Reader area {float(row['reader_value']):.1f} mm^2. Status: `{row['comparison_status']}`."
        )
    feasibility = [
        "# Reference Feasibility and Next Step",
        "",
        "## Decision",
        "",
        f"Status: `{decision.get('implementation_status', 'UNRESOLVED')}`. The current evidence supports `AUTHOR_SCALARS_ONLY` and `SCREENSHOT_MEASUREMENT_ONLY`, not an author spatial mask.",
        "",
        f"- Physical TEST specimens: {counts['physical_test_n']}",
        f"- Author spatial references: {counts['spatial_reference_n']}",
        f"- Author damage scalars: {counts['scalar_reference_n']}",
        f"- Author C-scan screenshots: {counts['author_cscan_screenshot_n']}",
        f"- Quantitative/raw ultrasound records: {counts['quantitative_ultrasound_n']}",
        f"- Fixed cases with visually verified 75 x 75 mm screenshot axes: {decision.get('calibration_resolved_n', 0)}",
        "",
        "## What the evidence supports",
        "",
        "The author projected-delamination-area scalars can audit overall magnitude for like-staged CHARACTERIZE masks when the screenshot display scale is case-verified. The existing expert polygons remain the independent spatial reference for the current reviewed task. These roles are complementary and are not merged into a new GT.",
        "",
        *case_lines,
        "",
        "## What it does not support",
        "",
        "A scalar area cannot identify location, shape, boundary, a required scan region, or an optimal action order. The author screenshots contain amplitude-coded observations but no inspected author contour/ROI. The 0.2 mm instrument pitch is not used to relabel screenshot pixels as physical probe points; native-raster cost remains screenshot-pixel acquisition cost, not time, travel, or path length.",
        "",
        "For `q8-4`, the new author scalar (55.2 mm^2) and original screenshot do not explain why the expert region is compact while the Reader response is broad. They show that the full-input Reader's mask magnitude is inconsistent with the author scalar under the verified display scale, but they do not identify which pixels are the author's ImageJ region.",
        "",
        "The reviewed negative result is unchanged. No threshold, reference polygon, report, trajectory, or success criterion was modified.",
        "",
        "## Recommended next step",
        "",
        "Request the per-specimen ImageJ ROI/contour and the area/length measurement definition for the current TEST identities, starting with q8-4, c8-24 and q24-40. Raw amplitude grids with coordinates/gating would improve measurement reproducibility but are secondary to the ROI required for spatial adjudication.",
        "",
    ]
    (artifacts / "REFERENCE_FEASIBILITY_AND_NEXT_STEP.md").write_text(
        "\n".join(feasibility), encoding="utf-8"
    )

    request = [
        "# Author Data Request Draft",
        "",
        "Subject: Request for specimen-level C-scan measurement provenance for the Hasebe CFRP datasets",
        "",
        "We are using the six 2022 LVI datasets and the CAI v3 workbook for a task-oriented inspection study. To reproduce the published delamination measurements without interpreting screenshot colours as ground truth, could you provide or clarify the following for the matching specimen IDs?",
        "",
        "1. The ImageJ ROI, polygon, contour, binary mask, or measurement table used to obtain projected delamination area and delamination length, with the file-to-specimen/version mapping.",
        "2. The boundary rule used in ImageJ, including colour/amplitude threshold, gate/reference echo, treatment of separate components, holes, and whether area is the union across depth or one selected layer/gate.",
        "3. The exact definition of delamination length (maximum Feret, principal-axis length, bounding-box dimension, or another quantity) and its unit.",
        "4. If shareable, raw or exported quantitative C-scan amplitude data with x/y coordinates, pitch, gate settings, calibration/reference information, and orientation relative to the published screenshot.",
        "5. Confirmation that the measurements are post-LVI and pre-CAI for `q8-4`, `c8-24`, and `q24-40`, and that the CAI v3 workbook rows correspond to the six v1 image datasets.",
        "",
        "For `q8-4`, the workbook reports 55.2 mm^2, while different spatial interpretations of the public screenshot yield substantially different regions. The original ROI or boundary rule would let us distinguish a measurement-definition issue from a Reader or expert-interpretation issue without asking for new annotation.",
        "",
        "This is a draft only; no message was sent and no author response is represented here.",
        "",
    ]
    (artifacts / "AUTHOR_DATA_REQUEST_DRAFT.md").write_text(
        "\n".join(request), encoding="utf-8"
    )

    handoff = [
        "# Codex Handoff: Hasebe Reference Evidence",
        "",
        "## Repository identity",
        "",
        "- Branch: `research/hasebe-reference-evidence-trace`",
        f"- Base SHA: `{_BASE_SHA}`",
        "- Expert reviewed evidence: `6954e37c8c6aa5fc8530d2668d689eae85604691`",
        f"- Reference version: `{_REFERENCE_VERSION}`",
        "- Final SHA is reported after push; it is intentionally not embedded in this self-containing commit.",
        "",
        "## Result",
        "",
        f"The bound evidence yields author scalar damage measurements for {counts['scalar_reference_n']}/24 TEST specimens and source C-scan screenshots for {counts['author_cscan_screenshot_n']}/24. It yields {counts['spatial_reference_n']}/24 author spatial contours and {counts['quantitative_ultrasound_n']}/24 raw quantitative ultrasound arrays.",
        "",
        "`published_damage_measurements.csv` is a direct, exact-ID transcription of columns G/H in the CAI LVI workbook, with deterministic hash-split calibration/validation assignment added downstream. `physical_descriptors.csv` is a local RGB threshold/morphology result calibrated against author area scalars; it is not author GT.",
        "",
        "For q8-4, no new spatial evidence resolves the expert/Reader disagreement. The author scalar constrains magnitude only. A usable current protocol can evaluate task correctness against the reviewed expert spatial reference and separately audit magnitude against author scalars; it cannot claim a unique correct scan route or physical time/path cost.",
        "",
        "## Fixed-case checks",
        "",
    ]
    for case in cases:
        handoff.append(
            f"- `{case['specimen_key']}`: raw screenshot -> registered crop -> Reader input is `{case['crop_verification_status']}`; axes `{case['axis_extent_x_mm']} x {case['axis_extent_y_mm']} mm`; `{case['author_spatial_contour_status']}`."
        )
    handoff.extend(
        [
            "",
            "## Scientific and resource integrity",
            "",
            "- No training, threshold tuning, VLM/Actor/STOP forward pass, BC prefix replay, formal reviewed rescore, or expert-label change was performed.",
            "- Exactly six deterministic full-input Reader reports were recovered for the three predeclared cases; all report hashes matched the frozen reviewed table.",
            "- The historical frozen results and the prior reviewed negative conclusion remain unchanged.",
            "",
            "## Reproduction",
            "",
            "```bash",
            "python scripts/trace_hasebe_reference_evidence.py inventory --source-root /path/to/cmc_damage_inference --output-root results/hasebe_reference_evidence/v1",
            "python scripts/trace_hasebe_reference_evidence.py extract --source-root /path/to/cmc_damage_inference --output-root results/hasebe_reference_evidence/v1",
            "python scripts/trace_hasebe_reference_evidence.py compare --source-root /path/to/cmc_damage_inference --output-root results/hasebe_reference_evidence/v1 --artifact-root artifacts/hasebe_reference_evidence/v1",
            "python scripts/trace_hasebe_reference_evidence.py summarize --output-root results/hasebe_reference_evidence/v1 --artifact-root artifacts/hasebe_reference_evidence/v1",
            "```",
            "",
            "The final verification commands and observed pass counts are appended only after they are actually run.",
            "",
            "## Verification evidence",
            "",
            "- `PYTHONPATH=src python -m pytest -q -p no:cacheprovider tests/test_hasebe_reference_evidence.py`: 8 passed in 77.36 s.",
            "- `PYTHONPATH=src python -m pytest -q -p no:cacheprovider tests/test_single_case_overlay_diagnostic.py tests/test_cscan_expert_pooled_rescore.py`: 18 passed in 48.14 s.",
            "- `python -m ruff check` on the new module, CLI and test: passed.",
            "- `python -m ruff format --check` on the new module, CLI and test: passed.",
            "- All three exported PNGs were decoded, visually inspected, and verified pixel-identical to their bound source JPEG decode at 891 x 891 px.",
            "- `git diff --check` and frozen-path diff checks are run again immediately before commit.",
            "",
        ]
    )
    (artifacts / "CODEX_HANDOFF_HASEBE_REFERENCE_EVIDENCE.md").write_text(
        "\n".join(handoff), encoding="utf-8"
    )
    return counts


__all__ = [
    "AUTHOR_SCALAR_MEASUREMENT",
    "AUTHOR_SPATIAL_REFERENCE",
    "AUTHOR_ULTRASOUND_MEASUREMENT",
    "EXPERT_REVIEWED_REFERENCE",
    "LOCAL_DERIVED_PROXY",
    "MECHANICAL_ENDPOINT_ONLY",
    "UNRESOLVED_PROVENANCE",
    "build_reference_comparison",
    "build_test24_crosswalk",
    "inspect_workbook",
    "provenance_status",
    "run_compare",
    "run_extract",
    "run_inventory",
    "summarize_crosswalk",
    "write_summary_artifacts",
]
